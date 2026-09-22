"""
Lector de ebooks/textos a voz en español con PyQt6 + edge-tts.

Formatos soportados:
- TXT, MD, CSV, SRT, VTT
- HTML / HTM
- EPUB        requiere: ebooklib beautifulsoup4
- PDF         requiere: pypdf
- DOCX        requiere: python-docx

Instalación recomendada:
    python -m pip install PyQt6 PyQt6-Qt6 PyQt6-sip edge-tts beautifulsoup4 ebooklib pypdf python-docx

Ejecutar:
    python app.py

Nota:
- edge-tts usa voces neuronales online. Necesita conexión a Internet.
- El audio se guarda como MP3.
"""

from __future__ import annotations

import asyncio
import html
import os
import re
import sys
import traceback
import tempfile
import uuid
from pathlib import Path
from typing import Callable, Dict, List, Optional

from PyQt6.QtCore import QObject, QThread, QUrl, Qt, pyqtSignal
from PyQt6.QtMultimedia import QAudioOutput, QMediaPlayer
from PyQt6.QtWidgets import (
    QApplication,
    QComboBox,
    QFileDialog,
    QFormLayout,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QProgressBar,
    QSlider,
    QSpinBox,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

try:
    import edge_tts
except ImportError:  # Se maneja con mensaje claro en la interfaz
    edge_tts = None


# ----------------------------- Lectura de archivos -----------------------------

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".srt", ".vtt", ".log", ".json", ".xml"}
HTML_EXTENSIONS = {".html", ".htm"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | HTML_EXTENSIONS | {".epub", ".pdf", ".docx", ".rtf"}


def normalize_text(text: str) -> str:
    """Limpia espacios excesivos sin destruir los párrafos."""
    text = html.unescape(text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\t\x0b\x0c]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ ]{2,}", " ", text)
    return text.strip()


def read_text_file(path: Path) -> str:
    """Lee texto simple probando varias codificaciones comunes."""
    encodings = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]
    last_error: Optional[Exception] = None
    for enc in encodings:
        try:
            return path.read_text(encoding=enc, errors="strict")
        except Exception as exc:  # noqa: BLE001
            last_error = exc
    # Último intento tolerante
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"No pude leer el archivo como texto. Error: {last_error or exc}")


def read_html(path: Path) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise RuntimeError("Para leer HTML instala: python -m pip install beautifulsoup4") from exc
    raw = read_text_file(path)
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n")


def read_epub(path: Path) -> str:
    try:
        import ebooklib
        from ebooklib import epub
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise RuntimeError("Para leer EPUB instala: python -m pip install ebooklib beautifulsoup4") from exc

    book = epub.read_epub(str(path))
    parts: List[str] = []
    for item in book.get_items():
        if item.get_type() == ebooklib.ITEM_DOCUMENT:
            soup = BeautifulSoup(item.get_content(), "html.parser")
            for tag in soup(["script", "style", "noscript"]):
                tag.decompose()
            text = soup.get_text("\n")
            if text.strip():
                parts.append(text)
    return "\n\n".join(parts)


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Para leer PDF instala: python -m pip install pypdf") from exc

    reader = PdfReader(str(path))
    pages: List[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            text = ""
        if text.strip():
            pages.append(f"\n\n--- Página {i} ---\n{text}")
    if not pages:
        raise RuntimeError("No pude extraer texto del PDF. Si es escaneado, necesita OCR.")
    return "".join(pages)


def read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("Para leer DOCX instala: python -m pip install python-docx") from exc

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)


def read_rtf(path: Path) -> str:
    """Extracción básica de RTF. Para documentos complejos, conviene convertir a TXT/DOCX."""
    raw = read_text_file(path)
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ").replace("\\", " ")
    return raw


def extract_text_from_file(file_path: str) -> str:
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise RuntimeError(f"Formato no soportado: {ext}")

    if ext in TEXT_EXTENSIONS:
        text = read_text_file(path)
    elif ext in HTML_EXTENSIONS:
        text = read_html(path)
    elif ext == ".epub":
        text = read_epub(path)
    elif ext == ".pdf":
        text = read_pdf(path)
    elif ext == ".docx":
        text = read_docx(path)
    elif ext == ".rtf":
        text = read_rtf(path)
    else:
        raise RuntimeError(f"Formato no soportado: {ext}")

    text = normalize_text(text)
    if not text:
        raise RuntimeError("El archivo no contiene texto legible.")
    return text


# ----------------------------- Síntesis de voz -----------------------------


def split_text_for_tts(text: str, max_chars: int = 3500) -> List[str]:
    """Divide el texto en fragmentos seguros para TTS, respetando párrafos y frases."""
    text = normalize_text(text)
    if not text:
        return []

    chunks: List[str] = []
    current = ""

    def flush() -> None:
        nonlocal current
        if current.strip():
            chunks.append(current.strip())
            current = ""

    paragraphs = re.split(r"\n\s*\n", text)
    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if not paragraph:
            continue

        if len(paragraph) > max_chars:
            # Divide por frases si el párrafo es enorme.
            sentences = re.split(r"(?<=[.!?¿¡;:])\s+", paragraph)
        else:
            sentences = [paragraph]

        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue

            if len(sentence) > max_chars:
                # Corte duro para frases extremadamente largas.
                flush()
                for i in range(0, len(sentence), max_chars):
                    chunks.append(sentence[i : i + max_chars].strip())
                continue

            candidate = f"{current}\n\n{sentence}" if current else sentence
            if len(candidate) <= max_chars:
                current = candidate
            else:
                flush()
                current = sentence

    flush()
    return chunks


async def synthesize_to_mp3(
    *,
    text: str,
    output_path: str,
    voice: str,
    rate_percent: int,
    volume_percent: int,
    pitch_hz: int,
    chunk_chars: int,
    progress_callback: Optional[Callable[[str, int], None]] = None,
) -> None:
    if edge_tts is None:
        raise RuntimeError("Falta edge-tts. Instala: python -m pip install edge-tts")

    chunks = split_text_for_tts(text, max_chars=chunk_chars)
    if not chunks:
        raise RuntimeError("No hay texto suficiente para convertir a voz.")

    rate = f"{rate_percent:+d}%"
    volume = f"{volume_percent:+d}%"
    pitch = f"{pitch_hz:+d}Hz"

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    # Escribimos todos los fragmentos en un único MP3.
    with output.open("wb") as audio_file:
        for index, chunk in enumerate(chunks, start=1):
            percent = int((index - 1) / len(chunks) * 100)
            if progress_callback:
                progress_callback(f"Convirtiendo fragmento {index}/{len(chunks)}...", percent)

            communicate = edge_tts.Communicate(
                text=chunk,
                voice=voice,
                rate=rate,
                volume=volume,
                pitch=pitch,
            )
            async for message in communicate.stream():
                if message["type"] == "audio":
                    audio_file.write(message["data"])

    if progress_callback:
        progress_callback("Conversión terminada.", 100)


# ----------------------------- Workers para no congelar la UI -----------------------------

class VoicesWorker(QObject):
    finished = pyqtSignal(list)
    error = pyqtSignal(str)

    def run(self) -> None:
        try:
            if edge_tts is None:
                raise RuntimeError("Falta edge-tts. Instala: python -m pip install edge-tts")
            voices = asyncio.run(edge_tts.list_voices())
            self.finished.emit(voices)
        except Exception as exc:  # noqa: BLE001
            self.error.emit(str(exc))


class TTSWorker(QObject):
    status = pyqtSignal(str)
    progress = pyqtSignal(int)
    finished = pyqtSignal(str)
    error = pyqtSignal(str)

    def __init__(
        self,
        *,
        text: str,
        output_path: str,
        voice: str,
        rate_percent: int,
        volume_percent: int,
        pitch_hz: int,
        chunk_chars: int,
    ) -> None:
        super().__init__()
        self.text = text
        self.output_path = output_path
        self.voice = voice
        self.rate_percent = rate_percent
        self.volume_percent = volume_percent
        self.pitch_hz = pitch_hz
        self.chunk_chars = chunk_chars

    def _progress_callback(self, message: str, percent: int) -> None:
        self.status.emit(message)
        self.progress.emit(percent)

    def run(self) -> None:
        try:
            asyncio.run(
                synthesize_to_mp3(
                    text=self.text,
                    output_path=self.output_path,
                    voice=self.voice,
                    rate_percent=self.rate_percent,
                    volume_percent=self.volume_percent,
                    pitch_hz=self.pitch_hz,
                    chunk_chars=self.chunk_chars,
                    progress_callback=self._progress_callback,
                )
            )
            self.finished.emit(self.output_path)
        except Exception as exc:  # noqa: BLE001
            details = traceback.format_exc()
            self.error.emit(f"{exc}\n\nDetalles técnicos:\n{details}")


# ----------------------------- Ventana principal -----------------------------

class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Lector de ebook/texto a voz - Español neural - PLAY texto escrito")
        self.resize(980, 760)

        self.current_file: Optional[str] = None
        self.current_output: Optional[str] = None
        self.all_voices: List[Dict[str, str]] = []
        self.auto_play_after_tts = False
        self.show_popup_after_tts = True

        self.player = QMediaPlayer(self)
        self.audio_output = QAudioOutput(self)
        self.player.setAudioOutput(self.audio_output)
        self.audio_output.setVolume(0.85)

        self._voices_thread: Optional[QThread] = None
        self._voices_worker: Optional[VoicesWorker] = None
        self._tts_thread: Optional[QThread] = None
        self._tts_worker: Optional[TTSWorker] = None

        self._build_ui()
        self._load_fallback_voices()
        self.load_voices_online()

    def _build_ui(self) -> None:
        root = QWidget()
        main_layout = QVBoxLayout(root)

        title = QLabel("Convertidor de ebooks y textos a voz agradable en español")
        title.setStyleSheet("font-size: 20px; font-weight: 700;")
        main_layout.addWidget(title)

        # Archivo
        file_box = QGroupBox("1) Cargar archivo")
        file_layout = QVBoxLayout(file_box)
        row_file = QHBoxLayout()
        self.file_line = QLineEdit()
        self.file_line.setReadOnly(True)
        self.file_line.setPlaceholderText("Selecciona un TXT, EPUB, PDF, DOCX, HTML, MD, etc.")
        btn_open = QPushButton("Abrir archivo")
        btn_open.clicked.connect(self.open_file)
        row_file.addWidget(self.file_line)
        row_file.addWidget(btn_open)
        file_layout.addLayout(row_file)
        self.file_info = QLabel("Formatos: TXT, MD, EPUB, PDF, DOCX, HTML, CSV, SRT, VTT, RTF")
        self.file_info.setStyleSheet("color: #555;")
        file_layout.addWidget(self.file_info)
        main_layout.addWidget(file_box)

        # Texto
        text_box = QGroupBox("2) Vista previa / texto editable")
        text_layout = QVBoxLayout(text_box)
        self.text_edit = QTextEdit()
        self.text_edit.setPlaceholderText("Escribe o pega aquí el texto que quieres probar y presiona el botón grande: ▶ PLAY AHORA.")
        text_layout.addWidget(self.text_edit)

        ayuda_play = QLabel("Para probar sin guardar archivo, escribe texto arriba y toca este botón:")
        ayuda_play.setStyleSheet("font-weight: 600;")
        text_layout.addWidget(ayuda_play)

        preview_row = QHBoxLayout()
        self.btn_play_text = QPushButton("▶ PLAY AHORA: leer el texto escrito en esta caja")
        self.btn_play_text.setMinimumHeight(48)
        self.btn_play_text.setToolTip("Lee en voz alta lo que escribiste o pegaste aquí. No tienes que convertir ni guardar MP3 primero.")
        self.btn_play_text.setStyleSheet("font-size: 16px; font-weight: bold; padding: 10px;")
        self.btn_play_text.clicked.connect(self.preview_written_text)
        preview_row.addWidget(self.btn_play_text)
        preview_row.addStretch(1)
        text_layout.addLayout(preview_row)

        self.char_label = QLabel("Caracteres: 0")
        self.text_edit.textChanged.connect(self.update_char_count)
        text_layout.addWidget(self.char_label)
        main_layout.addWidget(text_box, stretch=1)

        # Voz
        voice_box = QGroupBox("3) Voz neural en español")
        voice_layout = QFormLayout(voice_box)

        self.gender_combo = QComboBox()
        self.gender_combo.addItems(["Mujer", "Hombre", "Todas"])
        self.gender_combo.currentIndexChanged.connect(self.refresh_voice_combo)

        self.locale_combo = QComboBox()
        self.locale_combo.addItem("Todos los acentos", "ALL")
        self.locale_combo.currentIndexChanged.connect(self.refresh_voice_combo)

        self.voice_combo = QComboBox()

        btn_reload_voices = QPushButton("Actualizar voces online")
        btn_reload_voices.clicked.connect(self.load_voices_online)

        voice_layout.addRow("Género:", self.gender_combo)
        voice_layout.addRow("Acento / país:", self.locale_combo)
        voice_layout.addRow("Voz:", self.voice_combo)
        voice_layout.addRow("", btn_reload_voices)
        main_layout.addWidget(voice_box)

        # Ajustes
        settings_box = QGroupBox("4) Ajustes de lectura")
        settings_layout = QFormLayout(settings_box)

        self.rate_slider = QSlider(Qt.Orientation.Horizontal)
        self.rate_slider.setRange(-50, 50)
        self.rate_slider.setValue(0)
        self.rate_label = QLabel("0%")
        self.rate_slider.valueChanged.connect(lambda v: self.rate_label.setText(f"{v:+d}%"))
        rate_row = QHBoxLayout()
        rate_row.addWidget(self.rate_slider)
        rate_row.addWidget(self.rate_label)

        self.volume_slider = QSlider(Qt.Orientation.Horizontal)
        self.volume_slider.setRange(-50, 50)
        self.volume_slider.setValue(0)
        self.volume_label = QLabel("0%")
        self.volume_slider.valueChanged.connect(lambda v: self.volume_label.setText(f"{v:+d}%"))
        volume_row = QHBoxLayout()
        volume_row.addWidget(self.volume_slider)
        volume_row.addWidget(self.volume_label)

        self.pitch_slider = QSlider(Qt.Orientation.Horizontal)
        self.pitch_slider.setRange(-20, 20)
        self.pitch_slider.setValue(0)
        self.pitch_label = QLabel("0 Hz")
        self.pitch_slider.valueChanged.connect(lambda v: self.pitch_label.setText(f"{v:+d} Hz"))
        pitch_row = QHBoxLayout()
        pitch_row.addWidget(self.pitch_slider)
        pitch_row.addWidget(self.pitch_label)

        self.chunk_spin = QSpinBox()
        self.chunk_spin.setRange(1000, 8000)
        self.chunk_spin.setSingleStep(500)
        self.chunk_spin.setValue(3500)
        self.chunk_spin.setToolTip("Fragmentos más pequeños ayudan con libros largos.")

        settings_layout.addRow("Velocidad:", rate_row)
        settings_layout.addRow("Volumen de síntesis:", volume_row)
        settings_layout.addRow("Tono:", pitch_row)
        settings_layout.addRow("Tamaño de fragmento:", self.chunk_spin)
        main_layout.addWidget(settings_box)

        # Conversión y reproducción
        action_box = QGroupBox("5) Convertir y escuchar")
        action_layout = QVBoxLayout(action_box)
        row_actions = QHBoxLayout()
        self.btn_convert = QPushButton("Guardar como MP3")
        self.btn_convert.clicked.connect(self.convert_to_speech)
        self.btn_preview = QPushButton("▶ PLAY texto escrito")
        self.btn_preview.setToolTip("Lee en voz alta el texto escrito o pegado en el cuadro, sin pedir dónde guardar.")
        self.btn_preview.clicked.connect(self.preview_written_text)
        self.btn_play = QPushButton("Reproducir último MP3")
        self.btn_play.clicked.connect(self.play_audio)
        self.btn_pause = QPushButton("Pausar")
        self.btn_pause.clicked.connect(self.player.pause)
        self.btn_stop = QPushButton("Detener")
        self.btn_stop.clicked.connect(self.player.stop)
        row_actions.addWidget(self.btn_convert)
        row_actions.addWidget(self.btn_preview)
        row_actions.addWidget(self.btn_play)
        row_actions.addWidget(self.btn_pause)
        row_actions.addWidget(self.btn_stop)
        action_layout.addLayout(row_actions)

        self.progress = QProgressBar()
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.status_label = QLabel("Lista.")
        action_layout.addWidget(self.progress)
        action_layout.addWidget(self.status_label)
        main_layout.addWidget(action_box)

        self.setCentralWidget(root)

    def show_error(self, title: str, message: str) -> None:
        QMessageBox.critical(self, title, message)

    def show_info(self, title: str, message: str) -> None:
        QMessageBox.information(self, title, message)

    def update_char_count(self) -> None:
        count = len(self.text_edit.toPlainText())
        self.char_label.setText(f"Caracteres: {count:,}".replace(",", "."))

    def open_file(self) -> None:
        filters = (
            "Archivos de texto/ebook (*.txt *.md *.epub *.pdf *.docx *.html *.htm *.csv *.srt *.vtt *.rtf);;"
            "Todos los archivos (*.*)"
        )
        file_path, _ = QFileDialog.getOpenFileName(self, "Abrir archivo", "", filters)
        if not file_path:
            return
        try:
            self.status_label.setText("Leyendo archivo...")
            QApplication.processEvents()
            text = extract_text_from_file(file_path)
            self.current_file = file_path
            self.file_line.setText(file_path)
            self.text_edit.setPlainText(text)
            self.status_label.setText("Archivo cargado correctamente.")
        except Exception as exc:  # noqa: BLE001
            self.status_label.setText("No se pudo cargar el archivo.")
            self.show_error("Error al abrir archivo", str(exc))

    def _load_fallback_voices(self) -> None:
        """Voces conocidas para que la app tenga opciones aunque falle la lista online."""
        self.all_voices = [
            {"ShortName": "es-ES-ElviraNeural", "Gender": "Female", "Locale": "es-ES", "FriendlyName": "Microsoft Elvira Online (Natural) - Spanish (Spain)"},
            {"ShortName": "es-ES-AlvaroNeural", "Gender": "Male", "Locale": "es-ES", "FriendlyName": "Microsoft Alvaro Online (Natural) - Spanish (Spain)"},
            {"ShortName": "es-MX-DaliaNeural", "Gender": "Female", "Locale": "es-MX", "FriendlyName": "Microsoft Dalia Online (Natural) - Spanish (Mexico)"},
            {"ShortName": "es-MX-JorgeNeural", "Gender": "Male", "Locale": "es-MX", "FriendlyName": "Microsoft Jorge Online (Natural) - Spanish (Mexico)"},
            {"ShortName": "es-AR-ElenaNeural", "Gender": "Female", "Locale": "es-AR", "FriendlyName": "Microsoft Elena Online (Natural) - Spanish (Argentina)"},
            {"ShortName": "es-AR-TomasNeural", "Gender": "Male", "Locale": "es-AR", "FriendlyName": "Microsoft Tomas Online (Natural) - Spanish (Argentina)"},
            {"ShortName": "es-CO-SalomeNeural", "Gender": "Female", "Locale": "es-CO", "FriendlyName": "Microsoft Salome Online (Natural) - Spanish (Colombia)"},
            {"ShortName": "es-CO-GonzaloNeural", "Gender": "Male", "Locale": "es-CO", "FriendlyName": "Microsoft Gonzalo Online (Natural) - Spanish (Colombia)"},
        ]
        self.rebuild_locale_combo()
        self.refresh_voice_combo()

    def load_voices_online(self) -> None:
        if edge_tts is None:
            self.status_label.setText("Falta edge-tts. Instala con: python -m pip install edge-tts")
            return

        if self._voices_thread is not None and self._voices_thread.isRunning():
            return

        self.status_label.setText("Cargando voces neuronales en español...")
        self._voices_thread = QThread(self)
        self._voices_worker = VoicesWorker()
        self._voices_worker.moveToThread(self._voices_thread)
        self._voices_thread.started.connect(self._voices_worker.run)
        self._voices_worker.finished.connect(self.on_voices_loaded)
        self._voices_worker.error.connect(self.on_voices_error)
        self._voices_worker.finished.connect(self._voices_thread.quit)
        self._voices_worker.error.connect(self._voices_thread.quit)
        self._voices_thread.finished.connect(self._voices_worker.deleteLater)
        self._voices_thread.finished.connect(self._voices_thread.deleteLater)
        self._voices_thread.start()

    def on_voices_loaded(self, voices: List[Dict[str, str]]) -> None:
        spanish_voices = [v for v in voices if str(v.get("Locale", "")).lower().startswith("es-")]
        if spanish_voices:
            self.all_voices = sorted(
                spanish_voices,
                key=lambda v: (str(v.get("Locale", "")), str(v.get("Gender", "")), str(v.get("ShortName", ""))),
            )
            self.rebuild_locale_combo()
            self.refresh_voice_combo()
            self.status_label.setText(f"Voces cargadas: {len(spanish_voices)} voces en español.")
        else:
            self.status_label.setText("No encontré voces en español; usando voces predeterminadas.")

    def on_voices_error(self, message: str) -> None:
        self.status_label.setText("No pude actualizar voces online; se usan las predeterminadas.")
        # No mostramos popup para no interrumpir: la app sigue funcionando con presets.
        print("Error cargando voces:", message)

    def rebuild_locale_combo(self) -> None:
        current = self.locale_combo.currentData() if self.locale_combo.count() else "ALL"
        self.locale_combo.blockSignals(True)
        self.locale_combo.clear()
        self.locale_combo.addItem("Todos los acentos", "ALL")

        locales = sorted({str(v.get("Locale", "")) for v in self.all_voices if str(v.get("Locale", "")).startswith("es-")})
        for locale in locales:
            self.locale_combo.addItem(locale, locale)

        index = self.locale_combo.findData(current)
        self.locale_combo.setCurrentIndex(index if index >= 0 else 0)
        self.locale_combo.blockSignals(False)

    def refresh_voice_combo(self) -> None:
        wanted_gender = self.gender_combo.currentText()
        wanted_locale = self.locale_combo.currentData() if self.locale_combo.count() else "ALL"

        self.voice_combo.clear()
        for voice in self.all_voices:
            short_name = str(voice.get("ShortName", ""))
            gender = str(voice.get("Gender", ""))
            locale = str(voice.get("Locale", ""))
            friendly = str(voice.get("FriendlyName", short_name))

            if not locale.lower().startswith("es-"):
                continue
            if wanted_locale != "ALL" and locale != wanted_locale:
                continue
            if wanted_gender == "Mujer" and gender.lower() != "female":
                continue
            if wanted_gender == "Hombre" and gender.lower() != "male":
                continue

            etiqueta_genero = "Mujer" if gender.lower() == "female" else "Hombre" if gender.lower() == "male" else gender
            label = f"{short_name}  ·  {etiqueta_genero}  ·  {locale}"
            if friendly and friendly != short_name:
                label += f"  ·  {friendly}"
            self.voice_combo.addItem(label, short_name)

        if self.voice_combo.count() == 0:
            self.voice_combo.addItem("No hay voces para este filtro", "")

    def _start_tts_job(
        self,
        *,
        text: str,
        output_path: str,
        auto_play: bool,
        show_popup: bool,
    ) -> None:
        """Inicia una conversión TTS compartida por guardar MP3 y por prueba rápida."""
        self.auto_play_after_tts = auto_play
        self.show_popup_after_tts = show_popup
        self.btn_convert.setEnabled(False)
        self.btn_preview.setEnabled(False)
        if hasattr(self, "btn_play_text"):
            self.btn_play_text.setEnabled(False)
        self.progress.setValue(0)
        self.status_label.setText("Preparando conversión...")

        self._tts_thread = QThread(self)
        self._tts_worker = TTSWorker(
            text=text,
            output_path=output_path,
            voice=self.voice_combo.currentData(),
            rate_percent=self.rate_slider.value(),
            volume_percent=self.volume_slider.value(),
            pitch_hz=self.pitch_slider.value(),
            chunk_chars=self.chunk_spin.value(),
        )
        self._tts_worker.moveToThread(self._tts_thread)
        self._tts_thread.started.connect(self._tts_worker.run)
        self._tts_worker.status.connect(self.status_label.setText)
        self._tts_worker.progress.connect(self.progress.setValue)
        self._tts_worker.finished.connect(self.on_tts_finished)
        self._tts_worker.error.connect(self.on_tts_error)
        self._tts_worker.finished.connect(self._tts_thread.quit)
        self._tts_worker.error.connect(self._tts_thread.quit)
        self._tts_thread.finished.connect(self._tts_worker.deleteLater)
        self._tts_thread.finished.connect(self._tts_thread.deleteLater)
        self._tts_thread.finished.connect(lambda: self.btn_convert.setEnabled(True))
        self._tts_thread.finished.connect(lambda: self.btn_preview.setEnabled(True))
        self._tts_thread.finished.connect(lambda: self.btn_play_text.setEnabled(True) if hasattr(self, "btn_play_text") else None)
        self._tts_thread.start()

    def preview_written_text(self) -> None:
        """Convierte el texto escrito en la caja y lo reproduce de inmediato."""
        text = self.text_edit.toPlainText().strip()
        if not text:
            self.show_error("Sin texto", "Escribe o pega texto en el cuadro para probar la voz.")
            return

        voice = self.voice_combo.currentData()
        if not voice:
            self.show_error("Sin voz", "Selecciona una voz válida.")
            return

        # Para que la prueba sea rápida, usa solo el primer fragmento si el texto es muy largo.
        max_preview_chars = min(self.chunk_spin.value(), 2500)
        preview_text = text[:max_preview_chars].strip()
        if len(text) > max_preview_chars:
            preview_text += "\n\nVista previa recortada para prueba rápida."

        output_path = str(Path(tempfile.gettempdir()) / f"prueba_voz_{uuid.uuid4().hex}.mp3")
        self._start_tts_job(
            text=preview_text,
            output_path=output_path,
            auto_play=True,
            show_popup=False,
        )

    def convert_to_speech(self) -> None:
        text = self.text_edit.toPlainText().strip()
        if not text:
            self.show_error("Sin texto", "Carga un archivo o pega texto antes de convertir.")
            return

        voice = self.voice_combo.currentData()
        if not voice:
            self.show_error("Sin voz", "Selecciona una voz válida.")
            return

        default_name = "texto_a_voz.mp3"
        if self.current_file:
            default_name = f"{Path(self.current_file).stem}_voz.mp3"
        default_path = str(Path.home() / default_name)

        output_path, _ = QFileDialog.getSaveFileName(
            self,
            "Guardar audio MP3",
            default_path,
            "Audio MP3 (*.mp3)",
        )
        if not output_path:
            return
        if not output_path.lower().endswith(".mp3"):
            output_path += ".mp3"

        self._start_tts_job(
            text=text,
            output_path=output_path,
            auto_play=False,
            show_popup=True,
        )

    def on_tts_finished(self, output_path: str) -> None:
        self.current_output = output_path
        self.progress.setValue(100)
        self.player.setSource(QUrl.fromLocalFile(output_path))

        if self.auto_play_after_tts:
            self.status_label.setText("Prueba creada. Reproduciendo texto escrito...")
            self.player.play()
        else:
            self.status_label.setText(f"Audio creado: {output_path}")

        if self.show_popup_after_tts:
            self.show_info("Listo", f"Audio MP3 creado correctamente:\n{output_path}")

        self.auto_play_after_tts = False
        self.show_popup_after_tts = True

    def on_tts_error(self, message: str) -> None:
        self.progress.setValue(0)
        self.status_label.setText("Error durante la conversión.")
        if "Falta edge-tts" in message:
            self.show_error(
                "Falta instalar edge-tts",
                "Para que el botón PLAY y la conversión funcionen, instala edge-tts:\n\n"
                "python -m pip install edge-tts\n\n"
                "Después cierra y vuelve a abrir la app."
            )
        else:
            self.show_error("Error al convertir", message)

    def play_audio(self) -> None:
        if self.current_output and os.path.exists(self.current_output):
            self.player.setSource(QUrl.fromLocalFile(self.current_output))
            self.player.play()
            return

        # Permite abrir un MP3 ya generado.
        audio_path, _ = QFileDialog.getOpenFileName(self, "Abrir MP3", "", "Audio MP3 (*.mp3)")
        if audio_path:
            self.current_output = audio_path
            self.player.setSource(QUrl.fromLocalFile(audio_path))
            self.player.play()


def main() -> None:
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
